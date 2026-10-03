import json
import math
import re
import site
import sys
from datetime import datetime, date
from pathlib import Path


MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 10000
MAX_COLS_PER_SHEET = 200
MAX_TOTAL_CELLS = 250000


def emit(obj, code=0):
    print(json.dumps(obj, indent=2, ensure_ascii=False))
    raise SystemExit(code)


def load_openpyxl():
    errors = []

    def try_import():
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter, column_index_from_string
            from openpyxl.worksheet.table import Table, TableStyleInfo
            return {
                'Workbook': Workbook,
                'Alignment': Alignment,
                'Border': Border,
                'Font': Font,
                'PatternFill': PatternFill,
                'Side': Side,
                'get_column_letter': get_column_letter,
                'column_index_from_string': column_index_from_string,
                'Table': Table,
                'TableStyleInfo': TableStyleInfo,
            }
        except Exception as ex:
            errors.append('openpyxl: ' + str(ex))
            return None

    mod = try_import()
    if mod is not None:
        return mod, errors

    try:
        user_site = getattr(site, 'USER_SITE', None)
        if user_site:
            site.addsitedir(user_site)
    except Exception as ex:
        errors.append('user-site enable failed: ' + str(ex))

    mod = try_import()
    return mod, errors


def safe_name(stem):
    cleaned = re.sub(r'[^A-Za-z0-9._-]+', '_', str(stem or '')).strip('._-')
    return cleaned or 'workbook'


def safe_sheet_name(name, used):
    raw = str(name or 'Sheet').strip() or 'Sheet'
    raw = re.sub(r'[\\/*?:\[\]]+', ' ', raw).strip() or 'Sheet'
    base = raw[:31].strip() or 'Sheet'
    candidate = base
    i = 2
    while candidate.lower() in used:
        suffix = f'_{i}'
        candidate = (base[:31 - len(suffix)] + suffix).strip()
        i += 1
    used.add(candidate.lower())
    return candidate


def unique_path(folder, base_name):
    candidate = folder / base_name
    if not candidate.exists():
        return candidate
    stem = candidate.stem
    suffix = candidate.suffix
    for i in range(2, 1000):
        p = folder / f"{stem}_{i}{suffix}"
        if not p.exists():
            return p
    return folder / f"{stem}_latest{suffix}"


def normalize_filename(value):
    name = safe_name(str(value or 'workbook'))
    if not name.lower().endswith('.xlsx'):
        name += '.xlsx'
    return name


def col_to_index(ref, headers):
    if isinstance(ref, int):
        return ref if ref >= 1 else None
    s = str(ref or '').strip()
    if not s:
        return None
    if s.isdigit():
        n = int(s)
        return n if n >= 1 else None
    lowered = [str(h).strip().lower() for h in headers]
    if s.lower() in lowered:
        return lowered.index(s.lower()) + 1
    if re.fullmatch(r'[A-Za-z]{1,3}', s):
        try:
            from openpyxl.utils import column_index_from_string
            return column_index_from_string(s.upper())
        except Exception:
            return None
    return None



def truthy_flag(obj, keys, default=False):
    """Read flexible boolean flags from model-generated specs.

    Small local models do not always use the exact schema key.  This keeps
    formula permission explicit while accepting common aliases such as
    formulas_allowed, allow_formula, or formulas.  Nested options/settings
    dictionaries are also accepted.
    """
    def norm(v):
        if isinstance(v, bool):
            return v
        if isinstance(v, (int, float)):
            return v != 0
        s = str(v or '').strip().lower()
        if s in ('1', 'true', 'yes', 'y', 'on', 'allow', 'allowed', 'enable', 'enabled'):
            return True
        if s in ('0', 'false', 'no', 'n', 'off', 'deny', 'disabled', 'disable'):
            return False
        return bool(v)

    if not isinstance(obj, dict):
        return default
    for key in keys:
        if key in obj:
            return norm(obj.get(key))
    for parent_key in ('options', 'settings'):
        child = obj.get(parent_key)
        if isinstance(child, dict):
            for key in keys:
                if key in child:
                    return norm(child.get(key))
    return default


FORMULA_FLAG_KEYS = (
    'allow_formulas',
    'allow_formula',
    'formulas_allowed',
    'formula_allowed',
    'enable_formulas',
    'enable_formula',
    'formulas_enabled',
    'formula_enabled',
    'allow_excel_formulas',
    'formulas',
)

def coerce_value(value, allow_formulas=False):
    if isinstance(value, dict):
        # Friendly structured cell support:
        #   {"formula": "=A2+B2"}
        #   {"excel_formula": "SUM(A2:B2)"}
        # Formula objects still require allow_formulas=True; otherwise they are
        # protected as literal text just like formula-looking strings.
        formula = None
        for key in ('formula', 'excel_formula', 'xlsx_formula', 'cell_formula'):
            if key in value:
                formula = value.get(key)
                break
        if formula is not None:
            f = str(formula or '').strip()
            if f and not f.startswith('='):
                f = '=' + f
            if allow_formulas:
                return f[1:] if f.startswith("'=") else f
            return "'" + f if f.startswith(('=', '+', '-', '@')) else f

        # Common shape for richer future cells: {"value": 123, "format":"currency"}.
        for key in ('value', 'text', 'label'):
            if key in value:
                return coerce_value(value.get(key), allow_formulas)
        return ''

    if value is None:
        return ''
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        # JSON technically can't carry inf/nan, but defensive callers can.
        # openpyxl will happily write them and produce a workbook Excel flags
        # as "unreadable content".  Drop them to empty cells instead.
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return ''
        return value
    if isinstance(value, (datetime, date)):
        return value

    # Local models often serialize table numbers as JSON strings. Convert only
    # plain, safe numeric strings so formulas/text IDs are not damaged.
    # Examples converted: "2", "3", "10", "95.5", "0.92".
    # Examples preserved: "00123", "A-100", "2026-05-10", "=A2+B2".
    s = str(value)
    stripped = s.strip()

    if allow_formulas and stripped.startswith("'="):
        # Some model outputs pre-escape allowed formulas.  Only unescape an
        # explicit Excel formula when formulas were requested by the user.
        return stripped[1:]

    if not allow_formulas and stripped.startswith(('=', '+', '-', '@')):
        return "'" + s

    if re.fullmatch(r'-?(?:0|[1-9]\d*)(?:\.\d+)?', stripped):
        try:
            return float(stripped) if '.' in stripped else int(stripped)
        except Exception:
            pass

    return s

def rows_from_sheet(sheet_spec, headers):
    raw_rows = sheet_spec.get('rows', [])
    if raw_rows is None:
        raw_rows = []
    if not isinstance(raw_rows, list):
        raise ValueError('sheet rows must be a list')

    norm_headers = [_norm_label(h) for h in headers]

    rows = []
    for item in raw_rows:
        if isinstance(item, dict):
            # Exact-key match wins; fall back to normalized key match so
            # local-model output that varies in case/whitespace still maps
            # correctly. The normalized lookup is built lazily and only when
            # any exact-key lookup misses, so the common case stays cheap.
            norm_lookup = None
            row = []
            for h, nh in zip(headers, norm_headers):
                if h in item:
                    row.append(item[h])
                    continue
                if norm_lookup is None:
                    norm_lookup = {_norm_label(k): v for k, v in item.items()}
                if nh in norm_lookup:
                    row.append(norm_lookup[nh])
                else:
                    row.append('')
            rows.append(row)
        elif isinstance(item, list):
            rows.append(item)
        else:
            raise ValueError('each row must be an array or object')
    return rows


def parse_date_like(value):
    if isinstance(value, (datetime, date)):
        return value
    s = str(value or '').strip()
    if not s:
        return value
    # Common office-style date inputs. Keep the list intentionally small so
    # the helper does not guess ambiguous free-form text.
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m/%d/%y'):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            pass
    return value


def parse_percent_like(value):
    if isinstance(value, (int, float)):
        return value
    s = str(value or '').strip()
    if not s:
        return value
    if s.endswith('%'):
        try:
            return float(s[:-1].replace(',', '').strip()) / 100.0
        except Exception:
            return value
    try:
        return float(s.replace(',', ''))
    except Exception:
        return value


def parse_number_like(value, integer=False):
    if isinstance(value, (int, float)):
        return int(value) if integer else value
    s = str(value or '').strip()
    if not s:
        return value
    if s.startswith("'"):
        return value
    negative = False
    if s.startswith('(') and s.endswith(')'):
        negative = True
        s = s[1:-1].strip()
    s = s.replace('$', '').replace(',', '')
    try:
        n = float(s)
        if negative:
            n = -n
        return int(n) if integer else n
    except Exception:
        return value


def normalized_format_key(value):
    s = str(value or '').strip().lower()
    s = re.sub(r'[^a-z0-9%]+', '_', s).strip('_')
    aliases = {
        'currency': 'currency_columns',
        'money': 'currency_columns',
        'dollars': 'currency_columns',
        'dollar': 'currency_columns',
        'amount': 'currency_columns',
        'cost': 'currency_columns',
        'price': 'currency_columns',
        'number': 'number_columns',
        'numeric': 'number_columns',
        'decimal': 'number_columns',
        'float': 'number_columns',
        'integer': 'integer_columns',
        'int': 'integer_columns',
        'whole': 'integer_columns',
        'whole_number': 'integer_columns',
        'percent': 'percent_columns',
        'percentage': 'percent_columns',
        '%': 'percent_columns',
        'date': 'date_columns',
    }
    return aliases.get(s)


def add_format_ref(targets, key, ref):
    if not key or ref is None:
        return
    key = normalized_format_key(key) if not str(key).endswith('_columns') else key
    if key not in targets:
        return
    if isinstance(ref, str):
        ref = ref.strip()
    if ref == '':
        return
    if ref not in targets[key]:
        targets[key].append(ref)


def collect_column_format_refs(headers, sheet_spec):
    targets = {
        'currency_columns': [],
        'number_columns': [],
        'integer_columns': [],
        'percent_columns': [],
        'date_columns': [],
    }

    # Native schema: "percent_columns": ["Completion %"]
    for key in targets:
        raw = sheet_spec.get(key, []) or []
        if isinstance(raw, (str, int)):
            raw = [raw]
        if isinstance(raw, list):
            for ref in raw:
                add_format_ref(targets, key, ref)

    # Friendly aliases local models often produce:
    # "column_formats": {"Completion %": "percent", "Review Date": "date"}
    for map_key in ('column_formats', 'column_format', 'formats', 'format_columns'):
        raw = sheet_spec.get(map_key, None)
        if isinstance(raw, dict):
            for ref, fmt_name in raw.items():
                add_format_ref(targets, fmt_name, ref)
        elif isinstance(raw, list):
            for item in raw:
                if not isinstance(item, dict):
                    continue
                ref = item.get('column', item.get('header', item.get('name', item.get('ref'))))
                fmt_name = item.get('format', item.get('type', item.get('number_format')))
                add_format_ref(targets, fmt_name, ref)

    # Another common shape: "columns": [{"name":"Review Date", "format":"date"}]
    raw_columns = sheet_spec.get('columns', []) or []
    if isinstance(raw_columns, list):
        for item in raw_columns:
            if not isinstance(item, dict):
                continue
            ref = item.get('name', item.get('header', item.get('column')))
            fmt_name = item.get('format', item.get('type', item.get('number_format')))
            add_format_ref(targets, fmt_name, ref)

    # Safe inference from headers. This protects common natural-language cases
    # where the model builds the workbook data but forgets the format arrays.
    for h in headers:
        label = str(h or '').strip()
        lowered = label.lower()
        tokens = set(re.findall(r'[a-z0-9]+', lowered))
        if '%' in label or 'percent' in tokens or 'percentage' in tokens:
            add_format_ref(targets, 'percent', label)
        if 'date' in tokens or lowered.endswith(' date'):
            add_format_ref(targets, 'date', label)
        if {'cost', 'amount', 'price', 'budget', 'revenue', 'expense', 'total'} & tokens:
            # Avoid making generic "Total" integer/number columns currency unless
            # it clearly indicates money. "Total Cost" will be caught by cost.
            if {'cost', 'amount', 'price', 'budget', 'revenue', 'expense'} & tokens:
                add_format_ref(targets, 'currency', label)
        if {'count', 'qty', 'quantity', 'items'} & tokens or lowered.startswith('#'):
            add_format_ref(targets, 'integer', label)

    return targets


def apply_column_formats(ws, headers, sheet_spec, first_data_row, last_data_row):
    """Apply per-column number formats. Returns a dict of adjustment counts
    so the helper output can surface what was auto-normalized."""
    adjustments = {'percent_auto_decimal': 0}
    if last_data_row < first_data_row:
        return adjustments

    formats = {
        'currency_columns': '$#,##0.00',
        'number_columns': '#,##0.00',
        'integer_columns': '0',
        'percent_columns': '0.00%',
        'date_columns': 'yyyy-mm-dd',
    }
    targets = collect_column_format_refs(headers, sheet_spec)
    for key, fmt in formats.items():
        for ref in targets.get(key, []) or []:
            idx = col_to_index(ref, headers)
            if not idx:
                continue

            if key == 'percent_columns':
                # Pass 1: parse "92%" / "0.92" strings into numbers.
                for row in range(first_data_row, last_data_row + 1):
                    cell = ws.cell(row=row, column=idx)
                    cell.value = parse_percent_like(cell.value)
                # Pass 2: column-wide inference. If any value in the
                # column is numeric and outside [-1, 1], the column was
                # almost certainly authored in whole-percent form (e.g.
                # 92 meaning 92%, which 0.00% would otherwise render as
                # 9200%). Divide every numeric cell in the column by 100.
                whole_form = False
                for row in range(first_data_row, last_data_row + 1):
                    v = ws.cell(row=row, column=idx).value
                    if isinstance(v, bool):
                        continue
                    if isinstance(v, (int, float)) and (v > 1 or v < -1):
                        whole_form = True
                        break
                if whole_form:
                    for row in range(first_data_row, last_data_row + 1):
                        cell = ws.cell(row=row, column=idx)
                        v = cell.value
                        if isinstance(v, bool):
                            continue
                        if isinstance(v, (int, float)):
                            cell.value = v / 100.0
                            adjustments['percent_auto_decimal'] += 1
                for row in range(first_data_row, last_data_row + 1):
                    ws.cell(row=row, column=idx).number_format = fmt
            else:
                for row in range(first_data_row, last_data_row + 1):
                    cell = ws.cell(row=row, column=idx)
                    if key == 'date_columns':
                        cell.value = parse_date_like(cell.value)
                    elif key == 'integer_columns':
                        cell.value = parse_number_like(cell.value, integer=True)
                    elif key in ('number_columns', 'currency_columns'):
                        cell.value = parse_number_like(cell.value, integer=False)
                    cell.number_format = fmt

    # Optional precise cell formatting for mixed summary sheets, e.g.
    # "cell_formats": {"B5": "currency", "B2": "number"}
    named = {
        'currency': '$#,##0.00',
        'money': '$#,##0.00',
        'number': '#,##0.00',
        'integer': '0',
        'percent': '0.00%',
        'date': 'yyyy-mm-dd',
    }
    cell_formats = sheet_spec.get('cell_formats', {}) or {}
    if isinstance(cell_formats, dict):
        for addr, fmt_name in cell_formats.items():
            fmt_key = normalized_format_key(fmt_name)
            fmt = named.get(str(fmt_name).strip().lower()) or formats.get(fmt_key or '')
            if fmt:
                try:
                    cell = ws[str(addr)]
                    if fmt_key == 'date_columns':
                        cell.value = parse_date_like(cell.value)
                    elif fmt_key == 'percent_columns':
                        cell.value = parse_percent_like(cell.value)
                    elif fmt_key == 'integer_columns':
                        cell.value = parse_number_like(cell.value, integer=True)
                    elif fmt_key in ('number_columns', 'currency_columns'):
                        cell.value = parse_number_like(cell.value, integer=False)
                    cell.number_format = fmt
                except Exception:
                    pass
    return adjustments


def autosize_columns(ws, max_col):
    from openpyxl.utils import get_column_letter
    # Cells covered by a merged range (typically a long title spanning all
    # columns) shouldn't drive column widths -- the value lives in the
    # top-left cell only and stretching column A to fit the title produces
    # an awkward layout.
    merged_cells = set()
    for mr in ws.merged_cells.ranges:
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                merged_cells.add((r, c))

    for col_idx in range(1, max_col + 1):
        width = 10
        format_inflation = 0
        for row_idx in range(1, min(ws.max_row, 250) + 1):
            if (row_idx, col_idx) in merged_cells:
                continue
            cell = ws.cell(row=row_idx, column=col_idx)
            value = cell.value
            if value is not None:
                width = max(width, len(str(value)))
            fmt = (cell.number_format or '').strip()
            if fmt and fmt != 'General':
                # Display width exceeds raw len(str(value)) when the
                # format adds a currency symbol, thousands separators,
                # or a percent sign. Inflate by the worst-case
                # difference for the formats this helper emits.
                if '$' in fmt:
                    format_inflation = max(format_inflation, 4)
                elif '#,##' in fmt:
                    format_inflation = max(format_inflation, 3)
                elif '%' in fmt:
                    format_inflation = max(format_inflation, 2)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + format_inflation + 2, 10), 42)


def add_table_if_requested(ws, sheet_name, start_row, end_row, max_col, openpyxl_mod):
    if end_row <= start_row or max_col <= 0:
        return False
    get_column_letter = openpyxl_mod['get_column_letter']
    Table = openpyxl_mod['Table']
    TableStyleInfo = openpyxl_mod['TableStyleInfo']
    ref = f"A{start_row}:{get_column_letter(max_col)}{end_row}"
    display = 'Table_' + re.sub(r'[^A-Za-z0-9_]+', '_', sheet_name)
    display = display[:200] or 'Table1'
    tab = Table(displayName=display, ref=ref)
    style = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tab.tableStyleInfo = style
    ws.add_table(tab)
    return True


def _next_sheet_object_after_close(s, close_pos):
    i = close_pos + 1
    while i < len(s) and s[i].isspace():
        i += 1
    if i >= len(s) or s[i] != ',':
        return False
    i += 1
    while i < len(s) and s[i].isspace():
        i += 1
    if i >= len(s) or s[i] != '{':
        return False
    i += 1
    while i < len(s) and s[i].isspace():
        i += 1
    return s.startswith('"name"', i) or s.startswith("'name'", i)


def _scan_container_boundaries(s):
    """Yield (close_idx, open_idx, close_ch, open_ch) for every site where two
    JSON containers sit adjacent at structural depth -- i.e. ']' or '}'
    immediately followed (whitespace allowed) by '[' or '{'.

    Boundaries that fall inside a quoted string are skipped.  Escape state is
    tracked so a literal \\" inside a string does not flip our string-mode
    flag.
    """
    n = len(s)
    in_string = False
    escape = False
    i = 0
    while i < n:
        ch = s[i]
        if in_string:
            if escape:
                escape = False
            elif ch == '\\':
                escape = True
            elif ch == '"':
                in_string = False
            i += 1
            continue
        if ch == '"':
            in_string = True
            i += 1
            continue
        if ch == ']' or ch == '}':
            j = i + 1
            while j < n and s[j].isspace():
                j += 1
            if j < n and (s[j] == '[' or s[j] == '{'):
                yield (i, j, ch, s[j])
        i += 1


def _scan_trailing_commas(s):
    """Yield positions of structural ',' immediately followed (whitespace
    allowed) by ']' or '}'.  String-aware: commas inside strings are ignored.
    """
    n = len(s)
    in_string = False
    escape = False
    i = 0
    while i < n:
        ch = s[i]
        if in_string:
            if escape:
                escape = False
            elif ch == '\\':
                escape = True
            elif ch == '"':
                in_string = False
            i += 1
            continue
        if ch == '"':
            in_string = True
            i += 1
            continue
        if ch == ',':
            j = i + 1
            while j < n and s[j].isspace():
                j += 1
            if j < n and (s[j] == ']' or s[j] == '}'):
                yield i
        i += 1


def _scan_spurious_closers(s):
    """Yield positions of structural ']' or '}' that look spurious -- a
    closer sandwiched between a sibling boundary, i.e. immediately preceded
    by ']' or '}' AND immediately followed (whitespace, optional ',', more
    whitespace) by '[' or '{'.

    The model emits these when it loses track of nesting depth, e.g.
        ..."notes."]},["2026-05-01"...   <-- spurious '}' between two rows
    String-aware: positions inside string literals are skipped.
    """
    n = len(s)
    in_string = False
    escape = False
    i = 0
    while i < n:
        ch = s[i]
        if in_string:
            if escape:
                escape = False
            elif ch == '\\':
                escape = True
            elif ch == '"':
                in_string = False
            i += 1
            continue
        if ch == '"':
            in_string = True
            i += 1
            continue
        if ch == ']' or ch == '}':
            k = i - 1
            while k >= 0 and s[k].isspace():
                k -= 1
            if k < 0 or (s[k] != ']' and s[k] != '}'):
                i += 1
                continue
            j = i + 1
            while j < n and s[j].isspace():
                j += 1
            if j < n and s[j] == ',':
                j += 1
                while j < n and s[j].isspace():
                    j += 1
            if j < n and (s[j] == '[' or s[j] == '{'):
                yield i
        i += 1


def _repair_invalid_json_escapes(s):
    r"""Remove invalid backslash escapes inside JSON strings.

    Local/smaller models sometimes over-escape formula-like text, for example:
        "=HYPERLINK(\http://bad\, \click\)"

    JSON only allows these escapes inside a string: \" \\ \/ \b \f \n \r \t
    and \uXXXX.  When a backslash is followed by any other character, keep the
    following character and drop only the invalid backslash.  The scan is
    string-aware so structural JSON outside strings is left untouched.
    """
    s = str(s or '')
    out = []
    in_string = False
    i = 0
    repaired = False
    hexdigits = set('0123456789abcdefABCDEF')
    n = len(s)

    while i < n:
        ch = s[i]
        if not in_string:
            out.append(ch)
            if ch == '"':
                in_string = True
            i += 1
            continue

        if ch == '\\':
            if i + 1 >= n:
                out.append(ch)
                i += 1
                continue

            nxt = s[i + 1]
            if nxt in ('"', '\\', '/', 'b', 'f', 'n', 'r', 't'):
                out.append(ch)
                out.append(nxt)
                i += 2
                continue

            if nxt == 'u':
                seq = s[i + 2:i + 6]
                if len(seq) == 4 and all(c in hexdigits for c in seq):
                    out.append(s[i:i + 6])
                    i += 6
                    continue
                # Invalid \u escape. Preserve the intended literal text.
                out.append(nxt)
                i += 2
                repaired = True
                continue

            # Invalid JSON escape such as \h, \, or \).  Preserve the
            # intended literal character but remove the illegal backslash.
            out.append(nxt)
            i += 2
            repaired = True
            continue

        out.append(ch)
        if ch == '"':
            in_string = False
        i += 1

    repaired_text = ''.join(out)
    return repaired_text if repaired else s


def load_workbook_spec(raw_spec):
    """Load a workbook JSON spec with narrow repairs for local-model JSON slips.

    Repairs (each candidate is round-tripped through json.loads before being
    accepted; nothing is fixed blindly):
      * insert ',' between adjacent containers at structural depth -- }{, }[,
        ]{, ][ -- never inside a quoted string;
      * strip trailing ',' before ']' or '}';
      * drop a spurious ']' or '}' sandwiched between two sibling elements,
        e.g. ...notes."]},["next-row" -- never inside a quoted string;
      * insert the missing ']' that closes a rows array immediately before the
        next sheet object (legacy slip);
      * insert one missing ']' at the end if the spec ends with '}]}'.
    """
    raw_spec = str(raw_spec or '').strip()
    try:
        return json.loads(raw_spec), None
    except Exception as first_ex:
        first_error = str(first_ex)

    candidates = []

    def add(cand, note):
        for existing, _ in candidates:
            if existing == cand:
                return
        candidates.append((cand, note))

    invalid_escape_repaired = _repair_invalid_json_escapes(raw_spec)
    if invalid_escape_repaired != raw_spec:
        add(invalid_escape_repaired, 'removed invalid backslash escapes inside JSON strings')

    # 1. Missing comma between two adjacent containers.  This is the screenshot
    #    bug: small models occasionally emit two sheet objects (or two row
    #    arrays) back-to-back without the separator.
    for close_idx, _open_idx, close_ch, open_ch in _scan_container_boundaries(raw_spec):
        cand = raw_spec[:close_idx + 1] + ',' + raw_spec[close_idx + 1:]
        add(cand, 'inserted a missing comma between ' + close_ch + ' and ' + open_ch)

    # 2. Trailing comma before a closing bracket.
    for pos in _scan_trailing_commas(raw_spec):
        cand = raw_spec[:pos] + raw_spec[pos + 1:]
        add(cand, 'removed a trailing comma before a closing bracket')

    # 3. Spurious ']' or '}' between two sibling array/object elements.
    #    The model emits these when it loses track of nesting depth, e.g.
    #    ..."notes."]},["2026-05-01"...   <-- the extra '}' is spurious.
    for pos in _scan_spurious_closers(raw_spec):
        cand = raw_spec[:pos] + raw_spec[pos + 1:]
        add(cand, "removed a spurious '" + raw_spec[pos] + "' between two sibling elements")

    # 4. Legacy regex repair: missing rows-array ']' before the next sheet.
    repaired = re.sub(r'(\]\s*)(\}\s*,\s*\{\s*["\']name["\']\s*:)', r'\1]\2', raw_spec)
    if repaired != raw_spec:
        add(repaired, 'added a missing rows-array close bracket before the next sheet object')

    # 5. Legacy structural repair: insert one missing ']' before any
    #    sheet-object close that is immediately followed by the next sheet
    #    object.  Covers the same slip when the regex form does not match
    #    exactly (e.g. single quotes, extra whitespace).
    for pos, ch in enumerate(raw_spec):
        if ch != '}':
            continue
        if not _next_sheet_object_after_close(raw_spec, pos):
            continue
        cand = raw_spec[:pos] + ']' + raw_spec[pos:]
        add(cand, 'added a missing rows-array close bracket before the next sheet object')

    # 6. Final-sheet repair: small models sometimes close the sheet object
    #    before closing the rows array at the end of a single-sheet workbook:
    #       "rows":[["2","3","=A2+B2"]}}]}
    #    Correct form is:
    #       "rows":[["2","3","=A2+B2"]]}]}
    #    This is a character substitution, not an insertion, because the model
    #    emitted '}' where the rows-list closing ']' belonged.
    if raw_spec.endswith(']}}]}'):
        add(raw_spec[:-5] + ']]}]}', 'replaced final mistaken sheet close with a rows-array close bracket')

    # 7. Legacy last-resort: a single missing ']' before the final sheets
    #    array close. Keep this after the substitution repair above so the
    #    safer exact pattern wins for single-sheet workbook specs.
    if raw_spec.endswith('}]}'):
        add(raw_spec[:-3] + ']}]}', 'added a missing final rows-array close bracket')

    for cand, note in candidates:
        try:
            return json.loads(cand), note
        except Exception:
            pass

    raise ValueError(first_error)


def format_parse_error(raw_spec, error_message):
    """Make a JSON parse failure actionable for the next agent iteration.

    The default `json` message reports 'char N' but the model never sees N; it
    just retries with the same slip.  This wraps the message with a short
    snippet of the surrounding text and a plain-English hint about the most
    likely fix.
    """
    msg = str(error_message or '')
    m = re.search(r'char (\d+)', msg)
    if not m:
        return 'Invalid JSON workbook spec: ' + msg
    pos = int(m.group(1))
    if not raw_spec:
        return 'Invalid JSON workbook spec: ' + msg + ' (no input received).'
    pos = max(0, min(pos, len(raw_spec) - 1))
    start = max(0, pos - 30)
    end = min(len(raw_spec), pos + 30)
    snippet = raw_spec[start:end].replace('\n', ' ').replace('\r', ' ')
    prefix = '...' if start > 0 else ''
    suffix = '...' if end < len(raw_spec) else ''
    head = msg.split(':')[0]
    hint = ''
    if 'Expecting' in msg and "','" in msg:
        hint = " Hint: likely a missing ',' between two values, rows, or sheet objects."
    elif 'Expecting value' in msg:
        hint = ' Hint: likely a trailing comma, an unquoted literal, or a stray separator.'
    elif 'Expecting property name' in msg:
        hint = ' Hint: likely a trailing comma before } or a key not wrapped in double quotes.'
    elif 'Unterminated string' in msg:
        hint = ' Hint: a string is missing its closing quote, or contains an unescaped " or backslash.'
    return ('Invalid JSON workbook spec at char ' + str(pos) + ': ' + head
            + '. Near: ' + prefix + snippet + suffix + '.' + hint)


def _norm_label(value):
    return re.sub(r'[^a-z0-9]+', ' ', str(value or '').strip().lower()).strip()


def _to_number(value):
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value or '').strip().replace('$', '').replace(',', '')
    if not s:
        return 0.0
    try:
        return float(s)
    except Exception:
        return 0.0


def auto_correct_common_summary(spec):
    """Best-effort correction for common Metric/Value summary sheets.

    The model should compute summaries, but small local models occasionally make
    arithmetic mistakes.  If the workbook has a data sheet with these common
    headers, and a Summary sheet with Metric/Value rows, correct matching
    values deterministically before writing the xlsx.
    """
    adjustments = []
    sheets = spec.get('sheets') if isinstance(spec, dict) else None
    if not isinstance(sheets, list):
        return adjustments

    data_sheet = None
    summary_sheet = None
    required = {'scheduled hours', 'actual hours', 'overtime hours', 'overtime cost', 'status'}

    for sh in sheets:
        if not isinstance(sh, dict):
            continue
        name_norm = _norm_label(sh.get('name', ''))
        headers = [_norm_label(h) for h in (sh.get('headers') or [])]
        if name_norm == 'summary':
            summary_sheet = sh
        if required.issubset(set(headers)):
            data_sheet = sh

    if not data_sheet or not summary_sheet:
        return adjustments

    headers = [_norm_label(h) for h in (data_sheet.get('headers') or [])]
    col = {h: i for i, h in enumerate(headers)}
    rows = data_sheet.get('rows') or []
    if not isinstance(rows, list):
        return adjustments

    def sum_col(name):
        idx = col.get(name)
        if idx is None:
            return 0.0
        total = 0.0
        for r in rows:
            if isinstance(r, list) and idx < len(r):
                total += _to_number(r[idx])
            elif isinstance(r, dict):
                # Fall back to original header spelling when object rows are used.
                for raw_h in data_sheet.get('headers') or []:
                    if _norm_label(raw_h) == name:
                        total += _to_number(r.get(raw_h))
                        break
        return round(total, 2)

    status_idx = col.get('status')
    status_counts = {}
    if status_idx is not None:
        for r in rows:
            status = ''
            if isinstance(r, list) and status_idx < len(r):
                status = r[status_idx]
            elif isinstance(r, dict):
                for raw_h in data_sheet.get('headers') or []:
                    if _norm_label(raw_h) == 'status':
                        status = r.get(raw_h, '')
                        break
            key = _norm_label(status)
            if key:
                status_counts[key] = status_counts.get(key, 0) + 1

    metric_values = {
        'total scheduled hours': sum_col('scheduled hours'),
        'total actual hours': sum_col('actual hours'),
        'total overtime hours': sum_col('overtime hours'),
        'total overtime cost': sum_col('overtime cost'),
    }

    summary_rows = summary_sheet.get('rows') or []
    if not isinstance(summary_rows, list):
        return adjustments

    for r in summary_rows:
        if not isinstance(r, list) or len(r) < 2:
            continue
        label = _norm_label(r[0])
        new_value = None
        if label in metric_values:
            new_value = metric_values[label]
        else:
            m = re.fullmatch(r'count of (.+?) shifts?', label)
            if m:
                status_key = _norm_label(m.group(1))
                new_value = status_counts.get(status_key, 0)
        if new_value is not None and r[1] != new_value:
            adjustments.append({'metric': str(r[0]), 'old_value': r[1], 'new_value': new_value})
            r[1] = new_value

    # Mixed Metric/Value summary sheet: only the overtime-cost value is currency.
    if isinstance(summary_sheet.get('headers'), list) and [_norm_label(h) for h in summary_sheet.get('headers', [])[:2]] == ['metric', 'value']:
        cell_formats = summary_sheet.get('cell_formats')
        if not isinstance(cell_formats, dict):
            cell_formats = {}
        for idx, r in enumerate(summary_rows, start=2 + (2 if summary_sheet.get('title') else 0)):
            if isinstance(r, list) and r and _norm_label(r[0]) == 'total overtime cost':
                cell_formats[f'B{idx}'] = 'currency'
        if cell_formats:
            summary_sheet['cell_formats'] = cell_formats

    return adjustments


def build_workbook(spec, out_dir, openpyxl_mod):
    Workbook = openpyxl_mod['Workbook']
    Font = openpyxl_mod['Font']
    PatternFill = openpyxl_mod['PatternFill']
    Border = openpyxl_mod['Border']
    Side = openpyxl_mod['Side']
    Alignment = openpyxl_mod['Alignment']

    if not isinstance(spec, dict):
        raise ValueError('top-level spec must be a JSON object')
    sheets = spec.get('sheets')
    if not isinstance(sheets, list) or not sheets:
        raise ValueError('spec.sheets must be a non-empty array')
    if len(sheets) > MAX_SHEETS:
        raise ValueError(f'too many sheets; max is {MAX_SHEETS}')

    filename = normalize_filename(spec.get('filename', 'workbook.xlsx'))
    out_dir.mkdir(parents=True, exist_ok=True)
    output_path = unique_path(out_dir, filename)

    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    header_font = Font(bold=True)
    title_font = Font(bold=True, size=14)
    fill = PatternFill('solid', fgColor='D9EAF7')
    title_fill = PatternFill('solid', fgColor='BDD7EE')
    thin = Side(style='thin', color='D9D9D9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_top = Alignment(wrap_text=True, vertical='top')

    used_names = set()
    total_cells = 0
    sheet_summaries = []
    workbook_allow_formulas = truthy_flag(spec, FORMULA_FLAG_KEYS, False)

    for sheet_spec in sheets:
        if not isinstance(sheet_spec, dict):
            raise ValueError('each sheet spec must be an object')
        sheet_name = safe_sheet_name(sheet_spec.get('name', 'Sheet'), used_names)
        headers = sheet_spec.get('headers', []) or []
        if not isinstance(headers, list):
            raise ValueError(f'{sheet_name}: headers must be a list')
        headers = [str(h) for h in headers]
        if len(headers) > MAX_COLS_PER_SHEET:
            raise ValueError(f'{sheet_name}: too many columns; max is {MAX_COLS_PER_SHEET}')

        rows = rows_from_sheet(sheet_spec, headers)
        if len(rows) > MAX_ROWS_PER_SHEET:
            raise ValueError(f'{sheet_name}: too many rows; max is {MAX_ROWS_PER_SHEET}')
        max_col = max(len(headers), max((len(r) for r in rows), default=0))
        total_cells += max(1, len(rows) + (1 if headers else 0)) * max(1, max_col)
        if total_cells > MAX_TOTAL_CELLS:
            raise ValueError(f'workbook too large; max total cells is {MAX_TOTAL_CELLS}')

        ws = wb.create_sheet(sheet_name)
        row_cursor = 1
        title = str(sheet_spec.get('title', '') or '').strip()
        if title:
            ws.cell(row=row_cursor, column=1, value=title)
            if max_col > 1:
                ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=max_col)
            cell = ws.cell(row=row_cursor, column=1)
            cell.font = title_font
            cell.fill = title_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[row_cursor].height = 22
            row_cursor += 2

        header_row = row_cursor if headers else None
        if headers:
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=row_cursor, column=c_idx, value=h)
                cell.font = header_font
                cell.fill = fill
                cell.border = border
                cell.alignment = wrap_top
            row_cursor += 1

        allow_formulas = truthy_flag(sheet_spec, FORMULA_FLAG_KEYS, workbook_allow_formulas)
        first_data_row = row_cursor
        for row in rows:
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=row_cursor, column=c_idx, value=coerce_value(value, allow_formulas))
                cell.border = border
                cell.alignment = wrap_top
            row_cursor += 1
        last_data_row = row_cursor - 1

        if headers:
            # Border blank cells inside the declared table width so the sheet stays tidy.
            for r in range(first_data_row, last_data_row + 1):
                for c in range(1, len(headers) + 1):
                    ws.cell(row=r, column=c).border = border
                    ws.cell(row=r, column=c).alignment = wrap_top

        fmt_adjustments = apply_column_formats(ws, headers, sheet_spec, first_data_row, last_data_row)

        if sheet_spec.get('freeze_top_row', True) and header_row:
            ws.freeze_panes = f'A{header_row + 1}'

        # Excel can report that it found unreadable content when a worksheet
        # has both a normal worksheet autoFilter and an Excel Table on the
        # same range.  Tables already include their own filter dropdowns, so
        # create the table first and only add a worksheet-level autoFilter when
        # no table was added for this sheet.
        table_added = False
        if bool(sheet_spec.get('table', False)) and headers:
            table_added = add_table_if_requested(ws, sheet_name, header_row, max(last_data_row, header_row), max(len(headers), max_col), openpyxl_mod)

        if (not table_added) and sheet_spec.get('auto_filter', True) and headers and last_data_row >= header_row:
            from openpyxl.utils import get_column_letter
            ws.auto_filter.ref = f'A{header_row}:{get_column_letter(max(len(headers), max_col))}{max(last_data_row, header_row)}'

        autosize_columns(ws, max(len(headers), max_col))
        sheet_summaries.append({
            'name': sheet_name,
            'rows_written': len(rows),
            'columns_written': max(len(headers), max_col),
            'has_title': bool(title),
            'table_added': table_added,
            'percent_auto_decimal': fmt_adjustments.get('percent_auto_decimal', 0),
        })

    wb.save(output_path)
    return output_path, sheet_summaries


def main():
    if len(sys.argv) != 3:
        emit({'ok': False, 'helper': 'xlsx_create_workbook', 'error': 'xlsx_create_workbook expects exactly two internal arguments: JSON workbook spec and fixed output directory'}, 2)

    raw_spec_arg = sys.argv[1]
    if raw_spec_arg.startswith('@'):
        try:
            raw_spec = Path(raw_spec_arg[1:]).read_text(encoding='utf-8')
        except Exception as ex:
            emit({'ok': False, 'helper': 'xlsx_create_workbook', 'error': 'Could not read JSON spec file: ' + str(ex)}, 3)
    else:
        raw_spec = raw_spec_arg

    out_dir = Path(sys.argv[2]).resolve()
    if len(raw_spec.encode('utf-8', errors='replace')) > 256 * 1024:
        emit({'ok': False, 'helper': 'xlsx_create_workbook', 'error': 'JSON spec is larger than the first-phase 256 KB safety cap'}, 3)

    try:
        spec, repair_note = load_workbook_spec(raw_spec)
    except Exception as ex:
        emit({'ok': False, 'helper': 'xlsx_create_workbook', 'error': format_parse_error(raw_spec, str(ex))}, 4)

    summary_adjustments = auto_correct_common_summary(spec)

    openpyxl_mod, import_errors = load_openpyxl()
    if openpyxl_mod is None:
        emit({
            'ok': False,
            'helper': 'xlsx_create_workbook',
            'error': 'The openpyxl Python package is required for xlsx_create_workbook. Install it with: py -3 -m pip install --user openpyxl',
            'details': import_errors[-6:],
        }, 8)

    try:
        output_path, sheet_summaries = build_workbook(spec, out_dir, openpyxl_mod)
    except Exception as ex:
        emit({'ok': False, 'helper': 'xlsx_create_workbook', 'error': 'Could not create workbook: ' + str(ex)}, 6)

    emit({
        'ok': True,
        'helper': 'xlsx_create_workbook',
        'output_path': str(output_path),
        'output_filename': output_path.name,
        'spreadsheets_dir': str(out_dir),
        'json_repaired': bool(repair_note),
        'repair_note': repair_note or '',
        'summary_adjustments': summary_adjustments,
        'sheets': sheet_summaries,
    })


if __name__ == '__main__':
    main()