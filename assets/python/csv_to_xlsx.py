
import csv
import json
import re
import site
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path


def emit(obj, code=0):
    print(json.dumps(obj, indent=2, ensure_ascii=False))
    raise SystemExit(code)


def load_openpyxl():
    errors = []

    def try_import():
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter
            return Workbook, Font, get_column_letter
        except Exception as ex:
            errors.append('openpyxl: ' + str(ex))
            return None, None, None

    Workbook, Font, get_column_letter = try_import()
    if Workbook is not None:
        return Workbook, Font, get_column_letter, errors

    # PythonRunner launches fixed helpers with -I. Re-add the user site
    # so `py -3 -m pip install --user openpyxl` works during local testing.
    try:
        user_site = getattr(site, 'USER_SITE', None)
        if user_site:
            site.addsitedir(user_site)
    except Exception as ex:
        errors.append('user-site enable failed: ' + str(ex))

    Workbook, Font, get_column_letter = try_import()
    return Workbook, Font, get_column_letter, errors


def is_relative_to(child, parent):
    try:
        child.relative_to(parent)
        return True
    except Exception:
        return False


def decode_text(path):
    raw = path.read_bytes()
    for enc in ('utf-8-sig', 'utf-8', 'cp1252'):
        try:
            return raw.decode(enc), enc, len(raw)
        except UnicodeDecodeError:
            pass
    return raw.decode('utf-8', errors='replace'), 'utf-8-replace', len(raw)


def safe_name(stem):
    cleaned = re.sub(r'[^A-Za-z0-9._-]+', '_', stem).strip('._-')
    return cleaned or 'workbook'


def unique_path(folder, base_name):
    candidate = folder / base_name
    if not candidate.exists():
        return candidate
    stem = candidate.stem
    suffix = candidate.suffix
    for i in range(2, 1000):
        p = folder / f"{stem}_{i}{suffix}"
        if not p.exists():
            return p
    return folder / f"{stem}_latest{suffix}"


def maybe_number(value):
    s = str(value).strip()
    if not s:
        return ''

    # Keep leading-zero identifiers as text.
    plain = s.replace(',', '')
    if re.fullmatch(r'[+-]?\d+', plain):
        digits = plain[1:] if plain[:1] in '+-' else plain
        if len(digits) > 1 and digits.startswith('0'):
            return s
        try:
            return int(plain)
        except Exception:
            return s

    if re.fullmatch(r'[+-]?(\d+\.\d*|\.\d+)([eE][+-]?\d+)?', plain) or re.fullmatch(r'[+-]?\d+[eE][+-]?\d+', plain):
        try:
            return float(plain)
        except Exception:
            return s

    return s


def main():
    if len(sys.argv) != 3:
        emit({'ok': False, 'helper': 'csv_to_xlsx', 'error': 'csv_to_xlsx expects exactly two internal arguments: input path and fixed output directory'}, 2)

    cwd = Path.cwd().resolve()
    raw_arg = sys.argv[1]
    out_dir = Path(sys.argv[2]).resolve()

    target = Path(raw_arg)
    if not target.is_absolute():
        target = cwd / target
    target = target.resolve()


    if target.suffix.lower() not in ('.csv', '.tsv'):
        emit({'ok': False, 'helper': 'csv_to_xlsx', 'error': 'Only .csv and .tsv files are supported by csv_to_xlsx', 'path': str(target)}, 4)

    if not target.exists() or not target.is_file():
        emit({'ok': False, 'helper': 'csv_to_xlsx', 'error': 'File not found', 'path': str(target)}, 5)

    max_file_bytes = 10 * 1024 * 1024
    size_bytes = target.stat().st_size
    if size_bytes > max_file_bytes:
        emit({'ok': False, 'helper': 'csv_to_xlsx', 'error': 'File is larger than the first-phase 10 MB safety cap', 'path': str(target), 'file_size_bytes': size_bytes}, 7)

    Workbook, Font, get_column_letter, import_errors = load_openpyxl()
    if Workbook is None:
        emit({
            'ok': False,
            'helper': 'csv_to_xlsx',
            'error': 'The openpyxl Python package is required for csv_to_xlsx. Install it with: py -3 -m pip install --user openpyxl',
            'details': import_errors[-6:],
        }, 8)

    text, encoding, _ = decode_text(target)
    sample_text = text[:65536]
    delimiter = '\t' if target.suffix.lower() == '.tsv' else ','
    dialect_name = 'default-tsv' if delimiter == '\t' else 'default-csv'
    warnings = []

    try:
        sniffed = csv.Sniffer().sniff(sample_text, delimiters=',\t;|')
        delimiter = sniffed.delimiter
        dialect_name = 'sniffed'
    except Exception:
        warnings.append('Could not confidently sniff delimiter; used extension/default delimiter.')

    rows = []
    max_rows = 100000
    max_excel_cols = 16384
    try:
        reader = csv.reader(text.splitlines(), delimiter=delimiter)
        for row in reader:
            if len(row) > max_excel_cols:
                emit({'ok': False, 'helper': 'csv_to_xlsx', 'error': 'CSV has more columns than Excel supports', 'columns': len(row), 'excel_max_columns': max_excel_cols}, 9)
            rows.append(row)
            if len(rows) >= max_rows:
                warnings.append('Stopped writing after 100000 rows.')
                break
    except Exception as ex:
        emit({'ok': False, 'helper': 'csv_to_xlsx', 'error': 'CSV parse failed: ' + str(ex), 'path': str(target)}, 6)

    out_dir.mkdir(parents=True, exist_ok=True)
    output_path = unique_path(out_dir, safe_name(target.stem) + '.xlsx')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Data'

    max_widths = []
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row, start=1):
            converted = maybe_number(value)
            ws.cell(row=r_idx, column=c_idx, value=converted)
            text_len = len(str(value)) if value is not None else 0
            if c_idx > len(max_widths):
                max_widths.append(text_len)
            else:
                max_widths[c_idx - 1] = max(max_widths[c_idx - 1], text_len)

    if rows:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = 'A2'
        if len(rows) > 1 and max_widths:
            ws.auto_filter.ref = ws.dimensions

    for idx, width in enumerate(max_widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 40)

    try:
        wb.save(output_path)
    except Exception as ex:
        emit({'ok': False, 'helper': 'csv_to_xlsx', 'error': 'Could not save workbook: ' + str(ex), 'output_path': str(output_path)}, 10)

    emit({
        'ok': True,
        'helper': 'csv_to_xlsx',
        'input_path': str(target),
        'output_path': str(output_path),
        'output_filename': output_path.name,
        'cwd': str(cwd),
        'spreadsheets_dir': str(out_dir),
        'file_size_bytes': size_bytes,
        'encoding': encoding,
        'delimiter': delimiter,
        'dialect': dialect_name,
        'rows_written': len(rows),
        'row_count_exact': len(rows) < max_rows,
        'column_count_max': max((len(r) for r in rows), default=0),
        'warnings': warnings,
    })


if __name__ == '__main__':
    main()
