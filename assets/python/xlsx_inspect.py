
import json
import site
import sys
from pathlib import Path


def emit(obj, code=0):
    print(json.dumps(obj, indent=2, ensure_ascii=False))
    raise SystemExit(code)


def load_openpyxl():
    errors = []

    def try_import():
        try:
            from openpyxl import load_workbook
            return load_workbook
        except Exception as ex:
            errors.append('openpyxl: ' + str(ex))
            return None

    loader = try_import()
    if loader is not None:
        return loader, errors

    # PythonRunner launches helpers with -I for a safer default. That
    # hides per-user site packages on many Windows installs. Re-add the
    # user site only for this fixed helper so `py -3 -m pip install --user
    # openpyxl` works during local testing.
    try:
        user_site = getattr(site, 'USER_SITE', None)
        if user_site:
            site.addsitedir(user_site)
    except Exception as ex:
        errors.append('user-site enable failed: ' + str(ex))

    loader = try_import()
    return loader, errors


def is_relative_to(child, parent):
    try:
        child.relative_to(parent)
        return True
    except Exception:
        return False


def cell_to_jsonable(value):
    # openpyxl gives Python primitives for most cell values, but
    # date/datetime/time and Decimal need string projection so the
    # result round-trips through json.dumps without surprises.
    if value is None:
        return ''
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def main():
    if len(sys.argv) != 2:
        emit({'ok': False, 'helper': 'xlsx_inspect', 'error': 'xlsx_inspect expects exactly one file path argument'}, 2)

    cwd = Path.cwd().resolve()
    raw_arg = sys.argv[1]
    target = Path(raw_arg)
    if not target.is_absolute():
        target = cwd / target
    target = target.resolve()


    if target.suffix.lower() != '.xlsx':
        emit({'ok': False, 'helper': 'xlsx_inspect', 'error': 'Only .xlsx files are supported by xlsx_inspect (no .xls, .xlsm with macros, or .xlsb)', 'path': str(target)}, 4)

    if not target.exists() or not target.is_file():
        emit({'ok': False, 'helper': 'xlsx_inspect', 'error': 'File not found', 'path': str(target)}, 5)

    max_file_bytes = 50 * 1024 * 1024
    size_bytes = target.stat().st_size
    if size_bytes > max_file_bytes:
        emit({'ok': False, 'helper': 'xlsx_inspect', 'error': 'Workbook is larger than the first-phase 50 MB safety cap', 'path': str(target), 'file_size_bytes': size_bytes}, 7)

    load_workbook, import_errors = load_openpyxl()
    if load_workbook is None:
        emit({
            'ok': False,
            'helper': 'xlsx_inspect',
            'error': 'The openpyxl Python package is required for xlsx_inspect. Install it with: py -3 -m pip install --user openpyxl',
            'details': import_errors[-6:],
        }, 8)

    warnings = []
    try:
        # read_only=True streams sheets without loading the whole
        # workbook into memory; data_only=True returns cached values
        # rather than formula text.  Both matter for inspection-of-
        # last-saved-state on real timesheet files.
        wb = load_workbook(filename=str(target), read_only=True, data_only=True)
    except Exception as ex:
        emit({'ok': False, 'helper': 'xlsx_inspect', 'error': 'Could not open workbook: ' + str(ex), 'path': str(target)}, 6)

    max_rows_to_scan_per_sheet = 100000
    max_sample_data_rows = 5
    max_columns_to_report = 200

    sheets_out = []
    for sheet_name in wb.sheetnames:
        try:
            ws = wb[sheet_name]
        except Exception as ex:
            warnings.append('Could not open sheet "' + sheet_name + '": ' + str(ex))
            continue

        rows_iter = ws.iter_rows(values_only=True)
        rows = []
        row_count = 0
        try:
            for row in rows_iter:
                row_count += 1
                if len(rows) < (max_sample_data_rows + 1):
                    rows.append(list(row))
                if row_count >= max_rows_to_scan_per_sheet:
                    warnings.append('Sheet "' + sheet_name + '": stopped counting at 100000 rows.')
                    break
        except Exception as ex:
            warnings.append('Sheet "' + sheet_name + '": row scan failed: ' + str(ex))

        header = list(rows[0]) if rows else []
        if len(header) > max_columns_to_report:
            header = header[:max_columns_to_report]
            warnings.append('Sheet "' + sheet_name + '": column count exceeds 200; truncated.')

        sample_rows = []
        for raw in rows[1:max_sample_data_rows + 1]:
            sample_rows.append([cell_to_jsonable(v) for v in raw[:max_columns_to_report]])

        column_count = max((len(r) for r in rows), default=0)

        sheets_out.append({
            'name': sheet_name,
            'rows_scanned': row_count,
            'row_count_exact': row_count < max_rows_to_scan_per_sheet,
            'column_count_sample': column_count,
            'columns': [cell_to_jsonable(h) for h in header],
            'sample_rows': sample_rows,
        })

    try:
        wb.close()
    except Exception:
        pass

    emit({
        'ok': True,
        'helper': 'xlsx_inspect',
        'path': str(target),
        'cwd': str(cwd),
        'file_size_bytes': size_bytes,
        'sheet_count': len(sheets_out),
        'sheets': sheets_out,
        'warnings': warnings,
    })


if __name__ == '__main__':
    main()
