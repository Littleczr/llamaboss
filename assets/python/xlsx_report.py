
import json
import re
import site
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path


def emit(obj, code=0):
    print(json.dumps(obj, indent=2, ensure_ascii=False))
    raise SystemExit(code)


def load_openpyxl():
    errors = []

    def try_import():
        try:
            from openpyxl import load_workbook
            return load_workbook
        except Exception as ex:
            errors.append('openpyxl: ' + str(ex))
            return None

    loader = try_import()
    if loader is not None:
        return loader, errors

    # PythonRunner launches helpers with -I for a safer default. That
    # hides per-user site packages on many Windows installs. Re-add the
    # user site only for this fixed helper so `py -3 -m pip install --user
    # openpyxl` works during local testing.
    try:
        user_site = getattr(site, 'USER_SITE', None)
        if user_site:
            site.addsitedir(user_site)
    except Exception as ex:
        errors.append('user-site enable failed: ' + str(ex))

    loader = try_import()
    return loader, errors


def is_relative_to(child, parent):
    try:
        child.relative_to(parent)
        return True
    except Exception:
        return False


def safe_name(stem):
    cleaned = re.sub(r'[^A-Za-z0-9._-]+', '_', stem).strip('._-')
    return cleaned or 'xlsx_report'


def unique_path(folder, base_name):
    candidate = folder / base_name
    if not candidate.exists():
        return candidate
    stem = candidate.stem
    suffix = candidate.suffix
    for i in range(2, 1000):
        p = folder / f"{stem}_{i}{suffix}"
        if not p.exists():
            return p
    return folder / f"{stem}_latest{suffix}"


def md_cell(value):
    s = '' if value is None else str(value)
    s = s.replace('\r', ' ').replace('\n', ' ')
    s = s.replace('|', '\\|')
    if len(s) > 80:
        s = s[:77] + '...'
    return s


def md_table(headers, rows):
    out = []
    out.append('| ' + ' | '.join(md_cell(h) for h in headers) + ' |')
    out.append('| ' + ' | '.join('---' for _ in headers) + ' |')
    for row in rows:
        out.append('| ' + ' | '.join(md_cell(v) for v in row) + ' |')
    return '\n'.join(out)


def as_float(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return None  # don't treat True/False as numeric
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return None
    try:
        return float(s.replace(',', ''))
    except Exception:
        return None


def cell_to_value(value):
    if value is None:
        return ''
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def main():
    if len(sys.argv) != 3:
        emit({'ok': False, 'helper': 'xlsx_report', 'error': 'xlsx_report expects exactly two internal arguments: input path and fixed output directory'}, 2)

    cwd = Path.cwd().resolve()
    raw_arg = sys.argv[1]
    docs_dir = Path(sys.argv[2]).resolve()

    target = Path(raw_arg)
    if not target.is_absolute():
        target = cwd / target
    target = target.resolve()


    if target.suffix.lower() != '.xlsx':
        emit({'ok': False, 'helper': 'xlsx_report', 'error': 'Only .xlsx files are supported by xlsx_report (no .xls, .xlsm with macros, or .xlsb)', 'path': str(target)}, 4)

    if not target.exists() or not target.is_file():
        emit({'ok': False, 'helper': 'xlsx_report', 'error': 'File not found', 'path': str(target)}, 5)

    max_file_bytes = 50 * 1024 * 1024
    size_bytes = target.stat().st_size
    if size_bytes > max_file_bytes:
        emit({'ok': False, 'helper': 'xlsx_report', 'error': 'Workbook is larger than the first-phase 50 MB safety cap', 'path': str(target), 'file_size_bytes': size_bytes}, 7)

    load_workbook, import_errors = load_openpyxl()
    if load_workbook is None:
        emit({
            'ok': False,
            'helper': 'xlsx_report',
            'error': 'The openpyxl Python package is required for xlsx_report. Install it with: py -3 -m pip install --user openpyxl',
            'details': import_errors[-6:],
        }, 8)

    warnings = []
    try:
        wb = load_workbook(filename=str(target), read_only=True, data_only=True)
    except Exception as ex:
        emit({'ok': False, 'helper': 'xlsx_report', 'error': 'Could not open workbook: ' + str(ex), 'path': str(target)}, 6)

    max_rows_per_sheet = 100000
    max_sample_data_rows = 10
    max_columns_to_report = 200

    summary_sheets = []  # JSON-side sheet records returned alongside the report
    report_blocks = []   # markdown sections per sheet

    for sheet_name in wb.sheetnames:
        try:
            ws = wb[sheet_name]
        except Exception as ex:
            warnings.append('Could not open sheet "' + sheet_name + '": ' + str(ex))
            continue

        rows = []
        try:
            for row in ws.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) >= max_rows_per_sheet:
                    warnings.append('Sheet "' + sheet_name + '": stopped scanning at 100000 rows.')
                    break
        except Exception as ex:
            warnings.append('Sheet "' + sheet_name + '": row scan failed: ' + str(ex))
            continue

        header = list(rows[0]) if rows else []
        data_rows = rows[1:] if len(rows) > 1 else []
        width = max((len(r) for r in rows), default=0)
        if width > max_columns_to_report:
            warnings.append('Sheet "' + sheet_name + '": column count exceeds 200; truncated for the report.')
            width = max_columns_to_report
        if rows and any(len(r) != max((len(r) for r in rows), default=0) for r in rows[:200]):
            warnings.append('Sheet "' + sheet_name + '": some sampled rows have inconsistent column counts.')

        columns = []
        for i in range(width):
            name = ''
            if i < len(header) and header[i] is not None and str(header[i]).strip():
                name = str(header[i]).strip()
            else:
                name = f'column_{i + 1}'
            columns.append(name)

        missing_counts = [0 for _ in range(width)]
        nonempty_counts = [0 for _ in range(width)]
        numeric_values = [[] for _ in range(width)]

        for row in data_rows:
            for i in range(width):
                value = row[i] if i < len(row) else None
                if value is None or (isinstance(value, str) and value.strip() == ''):
                    missing_counts[i] += 1
                    continue
                nonempty_counts[i] += 1
                f = as_float(value)
                if f is not None:
                    numeric_values[i].append(f)

        numeric_summary = []
        for i, vals in enumerate(numeric_values):
            if not vals:
                continue
            denom = max(1, nonempty_counts[i])
            if len(vals) / denom < 0.80:
                continue
            total = sum(vals)
            numeric_summary.append({
                'column': columns[i],
                'count': len(vals),
                'sum': round(total, 6),
                'min': round(min(vals), 6),
                'max': round(max(vals), 6),
                'average': round(total / len(vals), 6),
            })

        # Markdown section for this sheet
        block = []
        block.append(f'## Sheet: {sheet_name}')
        block.append('')
        block.append(md_table(['Metric', 'Value'], [
            ['Rows scanned', len(rows)],
            ['Data rows', len(data_rows)],
            ['Columns', width],
        ]))
        block.append('')
        block.append('### Columns')
        block.append('')
        column_rows = []
        for i, name in enumerate(columns):
            column_rows.append([i + 1, name, nonempty_counts[i], missing_counts[i]])
        block.append(md_table(['#', 'Column', 'Non-empty', 'Missing/blank'], column_rows))
        block.append('')

        if numeric_summary:
            block.append('### Numeric Summary')
            block.append('')
            block.append(md_table(['Column', 'Count', 'Sum', 'Min', 'Max', 'Average'], [
                [n['column'], n['count'], n['sum'], n['min'], n['max'], n['average']]
                for n in numeric_summary
            ]))
            block.append('')

        block.append('### Sample Rows')
        block.append('')
        sample = data_rows[:max_sample_data_rows]
        if sample and columns:
            padded = []
            for row in sample:
                padded.append([(cell_to_value(row[i]) if i < len(row) else '') for i in range(width)])
            block.append(md_table(columns, padded))
        else:
            block.append('_No data rows found._')
        block.append('')

        report_blocks.append('\n'.join(block))

        summary_sheets.append({
            'name': sheet_name,
            'rows_scanned': len(rows),
            'data_row_count': len(data_rows),
            'column_count': width,
            'columns': columns,
            'numeric_columns': numeric_summary,
        })

    try:
        wb.close()
    except Exception:
        pass

    docs_dir.mkdir(parents=True, exist_ok=True)
    report_name = safe_name(target.stem) + '_report.md'
    output_path = unique_path(docs_dir, report_name)

    header_lines = []
    header_lines.append(f'# XLSX Report: {target.name}')
    header_lines.append('')
    header_lines.append('Generated by LlamaBoss controlled Python backend (`xlsx_report`).')
    header_lines.append('')
    header_lines.append('## Summary')
    header_lines.append('')
    header_lines.append(md_table(['Metric', 'Value'], [
        ['Input file', str(target)],
        ['Output report', str(output_path)],
        ['File size', f'{size_bytes} bytes'],
        ['Sheets', len(summary_sheets)],
    ]))
    header_lines.append('')

    warnings_block = []
    warnings_block.append('## Warnings')
    warnings_block.append('')
    if warnings:
        for w in warnings:
            warnings_block.append(f'- {w}')
    else:
        warnings_block.append('- None')
    warnings_block.append('')

    output_text = '\n'.join(header_lines) + '\n'.join(report_blocks) + '\n' + '\n'.join(warnings_block)
    output_path.write_text(output_text, encoding='utf-8')

    emit({
        'ok': True,
        'helper': 'xlsx_report',
        'input_path': str(target),
        'output_path': str(output_path),
        'output_filename': output_path.name,
        'cwd': str(cwd),
        'documents_dir': str(docs_dir),
        'file_size_bytes': size_bytes,
        'sheet_count': len(summary_sheets),
        'sheets': summary_sheets,
        'warnings': warnings,
    })


if __name__ == '__main__':
    main()
